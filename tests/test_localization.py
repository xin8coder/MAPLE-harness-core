import io
import json
import re
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from liveopt_dsh.config import ServiceConfig
from liveopt_dsh.data_adapter import prepare_public_data
from liveopt_dsh.documents import prepare_uploaded_documents
from liveopt_dsh.explorer import solution_options
from liveopt_dsh.localization import localize_value, map_prepared_data, response_language, public_semantic_context
from liveopt_dsh.service import LiveOptService
from liveopt_dsh.store import StateStore


def prepared(store):
    return prepare_public_data(store, store.root, inline_tables={
        "订单": [{"订单编号": "001", "数量": 10, "客户": "张先生"}],
        "班组": [{"班组编号": "甲组", "人数": 3}],
    })


def mappings(summary):
    specs = {"订单": ("orders", {"订单编号": "id", "数量": "quantity", "客户": "customer"}),
             "班组": ("teams", {"班组编号": "id", "人数": "workers"})}
    return {item["name"]: {"name": specs[item["source_name"]][0], "display_name": item["source_name"],
                           "columns": {col["name"]: specs[item["source_name"]][1][col["source_name"]]
                                       for col in item["columns"]}, "id_column": "id"}
            for item in summary["schema_preview"]}


def test_lossless_mapping_immutable_source_and_scoped_reverse_labels(tmp_path):
    store = StateStore(tmp_path)
    raw = prepared(store)
    assert raw["mapping_required"]
    assert len({t["name"] for t in raw["tables"]}) == 2
    before = store.load_prepared_data(raw["prepared_data_id"])
    mapped = map_prepared_data(store, raw["prepared_data_id"], mappings(raw))
    assert store.load_prepared_data(raw["prepared_data_id"]) == before
    context = store.load_prepared_data(mapped["prepared_data_id"])["public_context"]
    assert context["tables"] == {"orders": [{"id": "001", "quantity": 10, "customer": "张先生"}],
                                 "teams": [{"id": "甲组", "workers": 3}]}
    assert "id" not in context["localization"]["labels"]
    assert localize_value(context["tables"], context["localization"]) == {
        "订单": [{"订单编号": "001", "数量": 10, "客户": "张先生"}],
        "班组": [{"班组编号": "甲组", "人数": 3}]}
    assert map_prepared_data(store, raw["prepared_data_id"], mappings(raw))["prepared_data_id"] == mapped["prepared_data_id"]


@pytest.mark.parametrize("failure", ["missing_table", "duplicate_table", "missing_column", "duplicate_column", "invalid_name", "invalid_id", "changed_unit"])
def test_rejects_invalid_llm_mappings(tmp_path, failure):
    store = StateStore(tmp_path)
    raw = prepared(store)
    spec = mappings(raw)
    first, second = list(spec)
    columns = list(spec[first]["columns"])
    labels = None
    if failure == "missing_table":
        spec.pop(second)
    elif failure == "duplicate_table":
        spec[second]["name"] = spec[first]["name"]
    elif failure == "missing_column":
        spec[first]["columns"].pop(columns[0])
    elif failure == "duplicate_column":
        spec[first]["columns"][columns[1]] = spec[first]["columns"][columns[0]]
    elif failure == "invalid_name":
        spec[first]["name"] = "../orders"
    elif failure == "invalid_id":
        spec[first]["id_column"] = "unknown"
    else:
        labels = {"quantity": "金额（美元）"}
    with pytest.raises(ValueError):
        map_prepared_data(store, raw["prepared_data_id"], spec, result_labels=labels)


def test_duplicate_xlsx_headers_rejected_without_losing_data(tmp_path):
    book = Workbook()
    book.active.append(["数量", "数量"])
    book.active.append([1, 2])
    stream = io.BytesIO()
    book.save(stream)
    store = StateStore(tmp_path)
    upload = store.save_upload(filename="数据.xlsx", payload=stream.getvalue(), media_type="application/octet-stream")
    assert upload["filename"] == "数据.xlsx"
    with pytest.raises(ValueError, match="duplicate Excel headers"):
        prepare_uploaded_documents(store, [upload["upload_id"]])


def test_language_uses_request_not_english_schema():
    assert response_language("帮我安排一下订单，先不要算", {}) == "zh"
    assert response_language("Please schedule these orders.", {"localization": {"language": "zh"}}) == "en"
    assert localize_value({"a": 1, "b": 2}, {"labels": {"a": "成本", "b": "成本"}}) == {"成本 (a)": 1, "成本 (b)": 2}


def test_complete_prose_and_units_reach_generator():
    rows = [{"事项": str(i), "说明": f"第{i}条完整业务规则"} for i in range(7)]
    payload = public_semantic_context({"tables": {"rules": rows, "orders": [{"id": "甲", "count": 100}]}})
    assert "第6条完整业务规则" in payload
    assert '"orders"' not in payload
    with pytest.raises(ValueError, match="truncated"):
        public_semantic_context({"documents": [{"text": "incomplete", "truncated": True}]})


def test_chinese_workbook_real_search_update_export_offline(tmp_path, monkeypatch):
    root = Path(__file__).resolve().parents[1] / "examples" / "garment_export"
    monkeypatch.syspath_prepend(str(root))
    from build_chinese_inputs import SCHEMA, VALUES, INITIAL, UPDATE, chinese_tables
    from build_inputs import build_inputs
    from offline_client import OfflineClient
    from validate_results import validate_candidates
    path = build_inputs(tmp_path, tables=chinese_tables(), filename="服装外贸排产数据.xlsx", texts=[])
    book = load_workbook(path, data_only=True)
    for sheet in book:
        assert not re.search(r"[A-Za-z]", sheet.title)
        for row in sheet.values:
            assert all(not re.search(r"[A-Za-z]", value) for value in row if isinstance(value, str))
    book.close()

    class ChineseClient(OfflineClient):
        def chat(self, messages, **kwargs):
            # Scripted provider only; operator binding, search and persistence are real.
            response = super().chat(messages, **kwargs)
            payload = json.loads(response["choices"][0]["message"]["content"])
            if "operations" in payload:
                for operation in payload["operations"]:
                    value = operation["match"]["id"]
                    operation["match"]["id"] = VALUES.get(value, value)
                response["choices"][0]["message"]["content"] = json.dumps(payload, ensure_ascii=False)
            return response

    service = LiveOptService(ServiceConfig(state_dir=tmp_path / "state", workspace_root=tmp_path),
                            client_factory=lambda **kwargs: ChineseClient())
    try:
        upload = service.store.save_upload(filename=path.name, payload=path.read_bytes(), media_type="application/octet-stream")
        raw = service.prepare_uploads([upload["upload_id"]])
        with pytest.raises(ValueError, match="liveopt_map_data"):
            service.submit_start(problem=INITIAL, prepared_data_id=raw["prepared_data_id"])
        spec = {}
        for item in raw["schema_preview"]:
            key = next(k for k, (name, _) in SCHEMA.items() if item["source_name"].endswith("_" + name))
            name, cols = SCHEMA[key]
            reverse = {v: k for k, v in cols.items()}
            spec[item["name"]] = {"name": "garment_export_" + key, "display_name": name,
                                  "columns": {col["name"]: reverse[col["source_name"]] for col in item["columns"]}}
        mapped = service.map_data(raw["prepared_data_id"], spec, result_labels={"cost_cny": "总成本（元）", "late_carton_hours": "逾期箱小时", "order_plan": "订单安排", "order_id": "订单编号"})
        started = service.submit_start(problem=INITIAL, prepared_data_id=mapped["prepared_data_id"], population_size=20, generations=3)
        assert service.wait_for_job(started["job_id"], timeout_seconds=20)["status"] == "succeeded"
        session = started["session_id"]
        service.export_results(session)
        initial_options = solution_options(service.store, session)
        assert validate_candidates(service.store.load_snapshot(session), initial_options)["status"] == "passed"
        update = service.submit_update(session_id=session, update=UPDATE)
        assert service.wait_for_job(update["job_id"], timeout_seconds=20)["status"] == "succeeded"
        exported = service.export_results(session)
        snapshot = service.store.load_snapshot(session)
        options = solution_options(service.store, session)
        assert validate_candidates(snapshot, options)["status"] == "passed"
        assert snapshot["public_context"]["tables"]["garment_export_teams"][1]["workers"] == 3
        assert snapshot["public_context"]["tables"]["garment_export_orders"][1]["due_hour"] == 4.5
        assert options["localization"] == initial_options["localization"]
        assert exported["objective_names"] == ["总成本（元）", "逾期箱小时"]
        assert {"solutions_localized.csv", "solutions_localized.json", "field_mapping.json"} <= {item["name"] for item in exported["downloads"]}
        text = service.store.resolve_artifact(f"sessions/{session}/exports/t001/solutions_localized.json").read_text()
        assert "英国二号" in text and "订单编号" in text
        assert solution_options(service.store, session, 0)["candidates"] == initial_options["candidates"]
        assert "worker_shifts" in exported["unmapped_result_fields"]
        relabeled = service.export_results(session, result_labels={"worker_shifts": "工人排班", "order_plan.start_hour": "开工时刻"})
        assert "worker_shifts" not in relabeled["unmapped_result_fields"]
        assert service.store.load_snapshot(session) == snapshot
        assert service.export_results(session)["localization"]["labels"]["worker_shifts"] == "工人排班"
        assert "工人排班" in solution_options(service.store, session)["candidates"][0]["localized_solution"]
        assert "worker_shifts" not in solution_options(service.store, session, 0)["localization"]["labels"]
        with pytest.raises(ValueError, match="unknown result"):
            service.export_results(session, result_labels={"imaginary_field": "捏造字段"})
    finally:
        service.close()
