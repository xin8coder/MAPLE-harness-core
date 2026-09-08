"""Uploaded-document ingestion for the isolated LiveOpt application."""

from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
from typing import Any

from .data_adapter import (
    _load_source,
    _materialize_column_group,
    _normalize_records,
    _records,
    _safe_name,
)
from .serialization import json_safe
from .store import StateStore


TABULAR_SUFFIXES = frozenset({".csv", ".tsv", ".json", ".jsonl", ".xlsx"})
TEXT_SUFFIXES = frozenset({".txt", ".md", ".markdown", ".pdf", ".docx"})
SUPPORTED_UPLOAD_SUFFIXES = TABULAR_SUFFIXES | TEXT_SUFFIXES


def prepare_uploaded_documents(
    store: StateStore,
    upload_ids: list[str],
    *,
    inline_tables: dict[str, list[dict[str, Any]]] | None = None,
    column_groups: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """Convert immutable uploads into one reusable public context."""

    ordered_ids = [str(value or "").strip() for value in upload_ids]
    if not ordered_ids or any(not value for value in ordered_ids):
        raise ValueError("provide at least one upload_id")
    if len(ordered_ids) > 12:
        raise ValueError("at most 12 documents may be prepared together")

    tables: dict[str, list[dict[str, Any]]] = {}
    dictionary: dict[str, Any] = {}
    documents: list[dict[str, Any]] = []
    sources: list[dict[str, Any]] = []
    total_text_budget = max(
        1_000,
        int(os.getenv("LIVEOPT_MCP_MAX_DOCUMENT_CHARS", "120000")),
    )
    remaining_text_budget = total_text_budget

    for upload_id in ordered_ids:
        metadata, path = store.resolve_upload(upload_id)
        filename = str(metadata["filename"])
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_UPLOAD_SUFFIXES:
            allowed = ", ".join(sorted(SUPPORTED_UPLOAD_SUFFIXES))
            raise ValueError(f"unsupported uploaded document {suffix!r}; use {allowed}")

        source_label = f"upload:{upload_id}:{filename}"
        if suffix == ".xlsx":
            loaded = _load_xlsx(path, filename)
            _add_tables(
                loaded,
                tables=tables,
                dictionary=dictionary,
                source=source_label,
                infer_scalars=False,
            )
            sources.append(
                {
                    "upload_id": upload_id,
                    "filename": filename,
                    "format": "xlsx",
                    "kind": "tables",
                }
            )
            continue
        if suffix in TABULAR_SUFFIXES:
            loaded = _load_source(path, requested_name=Path(filename).stem)
            _add_tables(
                loaded,
                tables=tables,
                dictionary=dictionary,
                source=source_label,
                infer_scalars=suffix in {".csv", ".tsv"},
            )
            sources.append(
                {
                    "upload_id": upload_id,
                    "filename": filename,
                    "format": suffix.removeprefix("."),
                    "kind": "tables",
                }
            )
            continue

        extracted = _extract_document_text(path, suffix)
        original_chars = len(extracted)
        per_document_limit = min(60_000, remaining_text_budget)
        retained = extracted[:per_document_limit]
        truncated = len(retained) < original_chars
        remaining_text_budget -= len(retained)
        documents.append(
            {
                "upload_id": upload_id,
                "filename": filename,
                "media_type": metadata["media_type"],
                "sha256": metadata["sha256"],
                "text": retained,
                "original_chars": original_chars,
                "retained_chars": len(retained),
                "truncated": truncated,
            }
        )
        sources.append(
            {
                "upload_id": upload_id,
                "filename": filename,
                "format": suffix.removeprefix("."),
                "kind": "document",
                "truncated": truncated,
            }
        )

    for raw_name, records in dict(inline_tables or {}).items():
        table_name = _unique_name(_safe_name(raw_name), tables)
        normalized, metadata = _normalize_records(
            _records(records), id_column="", infer_scalars=False
        )
        tables[table_name] = normalized
        dictionary[table_name] = {
            "source": "inline_public_data",
            "row_count": len(normalized),
            **metadata,
        }
        sources.append({"path": None, "format": "inline", "table": table_name})

    for group in list(column_groups or []):
        _materialize_column_group(tables, dictionary, sources, group)

    public_context = {
        "tables": tables,
        "data_dictionary": dictionary,
        "documents": documents,
        "prepared_sources": sources,
    }
    canonical = json.dumps(
        json_safe(public_context),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    data_id = hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:16]
    record = {
        "schema_version": 1,
        "prepared_data_id": data_id,
        "public_context": public_context,
    }
    artifact = store.save_prepared_data(data_id, record)
    return {
        "prepared_data_id": data_id,
        "artifact": str(artifact.relative_to(store.root)),
        "tables": [
            {
                "name": name,
                "rows": len(tables[name]),
                "columns": dictionary[name]["columns"],
                "id_column": dictionary[name]["id_column"],
            }
            for name in tables
        ],
        "documents": [
            {
                "upload_id": item["upload_id"],
                "filename": item["filename"],
                "retained_chars": item["retained_chars"],
                "truncated": item["truncated"],
            }
            for item in documents
        ],
    }


def _add_tables(
    loaded: dict[str, list[dict[str, Any]]],
    *,
    tables: dict[str, list[dict[str, Any]]],
    dictionary: dict[str, Any],
    source: str,
    infer_scalars: bool,
) -> None:
    for source_name, records in loaded.items():
        table_name = _unique_name(_safe_name(source_name), tables)
        normalized, metadata = _normalize_records(
            records,
            id_column="",
            infer_scalars=infer_scalars,
        )
        tables[table_name] = normalized
        dictionary[table_name] = {
            "source": source,
            "row_count": len(normalized),
            **metadata,
        }


def _unique_name(base: str, tables: dict[str, Any]) -> str:
    if base not in tables:
        return base
    counter = 2
    while f"{base}_{counter}" in tables:
        counter += 1
    return f"{base}_{counter}"


def _load_xlsx(path: Path, filename: str) -> dict[str, list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - installed application dependency
        raise RuntimeError("XLSX support requires openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    loaded: dict[str, list[dict[str, Any]]] = {}
    stem = Path(filename).stem
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            try:
                header = next(rows)
            except StopIteration:
                continue
            names = [str(value).strip() if value is not None else "" for value in header]
            if not any(names):
                continue
            names = [name or f"column_{index + 1}" for index, name in enumerate(names)]
            records = [
                {name: value for name, value in zip(names, row)}
                for row in rows
                if any(value is not None for value in row)
            ]
            loaded[f"{stem}_{worksheet.title}"] = records
    finally:
        workbook.close()
    if not loaded:
        raise ValueError(f"uploaded workbook has no tabular sheets: {filename}")
    return loaded


def _extract_document_text(path: Path, suffix: str) -> str:
    if suffix in {".txt", ".md", ".markdown"}:
        return path.read_text(encoding="utf-8-sig").strip()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - installed application dependency
            raise RuntimeError("PDF support requires pypdf") from exc
        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - installed application dependency
            raise RuntimeError("DOCX support requires python-docx") from exc
        document = Document(str(path))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(
                "\t".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
            )
        return "\n".join(value for value in paragraphs if value).strip()
    raise ValueError(f"unsupported document suffix: {suffix}")
