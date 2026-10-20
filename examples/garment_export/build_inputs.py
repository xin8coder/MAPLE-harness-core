"""Generate the public Excel workbook and equivalent conversational inputs."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from scenario import INITIAL_REQUEST, TABLES, UPDATE_REQUEST, natural_language_data


def build_inputs(output: Path, *, tables=None, filename="garment_export.xlsx", texts=None) -> Path:
    output.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, records in (TABLES if tables is None else tables).items():
        sheet = workbook.create_sheet(name)
        sheet.append(list(records[0]))
        for record in records:
            sheet.append(list(record.values()))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="087F78")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F0F6F4")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 22
        if name in {"rules", "业务规则"}:
            sheet.column_dimensions["B"].width = 94
            for row in range(2, sheet.max_row + 1):
                sheet.row_dimensions[row].height = 60
    file = output / filename
    workbook.save(file)
    for name, content in (texts if texts is not None else (
            ("initial_request.txt", INITIAL_REQUEST), ("update_1.txt", UPDATE_REQUEST),
            ("natural_language_data.txt", natural_language_data()))):
        (output / name).write_text(content + "\n", encoding="utf-8")
    return file


if __name__ == "__main__":
    print(build_inputs(Path(__file__).parent / "inputs"))
